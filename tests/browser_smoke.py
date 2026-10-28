"""Headless Edge acceptance checks with disposable SQLite and mocked Gemini HTTP.

Install optional QA dependency: pip install playwright
Uses the installed Microsoft Edge browser; no production database or real AI calls.
"""
import errno
import json
import os
from pathlib import Path
import re
import sys
import tempfile
import threading
import io
import urllib.error
from unittest.mock import patch

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT))
sys.path.insert(0,str(ROOT/'scratch/qa_audit_2026_10_27/browser_deps'))
os.environ.update(AI_CONFIG_SOURCE='process',AI_PROVIDER='gemini',GEMINI_API_KEY='qa_fake_key',
                  AI_LOCAL_FALLBACK_ENABLED='false',LOCAL_AI_MODEL_ENABLED='false',
                  GEMINI_MODEL='gemini-3.6-flash',GEMINI_MEETING_MODEL='gemini-3.6-flash',
                  SMTP_ENABLED='false',DEADLINE_WORKER_ENABLED='false',FLASK_DEBUG='false')
from playwright.sync_api import sync_playwright, expect


def main():
    from database import connection,init_db
    records=[]; errors=[]; calls=[]; network_mode={'denied':False}
    out=ROOT/'docs'; shots=out/'qa-screenshots'; shots.mkdir(parents=True,exist_ok=True)
    def record(name,details): records.append({'name':name,'passed':True,'details':details})
    def gemini_http(request, timeout=None):
        assert request.full_url == 'https://generativelanguage.googleapis.com/v1beta/models/gemini-3.6-flash:generateContent'
        payload = json.loads(request.data)
        calls.append('gemini-3.6-flash')
        if network_mode['denied']:
            denied = PermissionError(errno.EACCES, 'qa_transport_detail')
            denied.winerror = 10013
            raise urllib.error.URLError(denied)
        prompt = payload['contents'][0]['parts'][0]['text']
        if 'NGUỒN:\n' in prompt:
            ids = re.findall(r'^\[(\d+)\]', prompt.split('NGUỒN:\n')[-1], re.MULTILINE)
            text = json.dumps({'items':[{'id':int(uid), 'kinds':['discussion'], 'summary':'Trao đổi theo nguồn cuộc họp.'} for uid in ids]})
        else:
            text = 'Bản nháp đối chiếu với dữ liệu dự án được cung cấp.'
        return io.BytesIO(json.dumps({'candidates':[{'finishReason':'STOP',
            'content':{'parts':[{'text':text}]}}]}).encode('utf-8'))
    with tempfile.TemporaryDirectory() as directory:
        connection._USE_MYSQL=False
        connection._SQLITE_PATH=init_db.SQLITE_PATH=str(Path(directory)/'browser.sqlite3')
        init_db.init_sqlite_fallback()
        from app import app
        from werkzeug.serving import make_server
        server=make_server('0.0.0.0',0,app,threaded=True)
        port=server.server_port; base=f'http://127.0.0.1:{port}'
        worker=threading.Thread(target=server.serve_forever,daemon=True);worker.start()
        with patch('services.gemini_service.urllib.request.urlopen', side_effect=gemini_http),sync_playwright() as pw:
            browser=pw.chromium.launch(channel='msedge',headless=True)
            def context(width=1366,height=900):
                ctx=browser.new_context(viewport={'width':width,'height':height},locale='vi-VN',is_mobile=width<600,has_touch=width<1000)
                ctx.route('https://fonts.googleapis.com/**',lambda route:route.abort())
                ctx.route('https://fonts.gstatic.com/**',lambda route:route.abort())
                ctx.on('page',lambda page:page.on('pageerror',lambda error:errors.append(str(error))))
                return ctx
            def login(ctx,username):
                page=ctx.new_page();page.goto(base+'/login')
                page.locator('#login-username').fill(username);page.locator('#login-password').fill('Password123!')
                page.locator('#btn-submit-login').click();page.wait_for_url('**/'+('admin' if username=='admin' else 'dashboard'))
                page.wait_for_function('currentUser !== null')
                return page
            try:
                leader_ctx=context();leader=login(leader_ctx,'leader_long')
                member_ctx=context();member=login(member_ctx,'member_my')
                admin_ctx=context();admin=login(admin_ctx,'admin')
                record('Login: Admin, PM and member',{'all_pages':200})
                network_mode['denied']=True
                admin.get_by_role('button',name='Kiểm tra API',exact=True).click()
                expect(admin.locator('#admin-ai-status')).to_contain_text('AI_NETWORK_DENIED')
                network_mode['denied']=False
                admin.get_by_role('button',name='Kiểm tra API',exact=True).click()
                expect(admin.locator('#admin-ai-status')).to_contain_text('Đã kiểm tra tạo nội dung')
                record('Admin diagnostics distinguish network denied and generation reachable',{'error_code':'AI_NETWORK_DENIED','recovery':True,'quota_not_claimed':True})

                # Real modal interaction, Unicode rendering and initial staff selection.
                leader.get_by_role('button',name='+ Tạo dự án',exact=True).click()
                expect(leader.locator('.initial-member-option').filter(has_text='Lưu Trà My')).to_be_visible()
                leader.locator('#new-project-name').fill('Dự án tiếng Việt đồng bộ')
                leader.locator('#new-project-code').fill('BROWSER-QA')
                leader.locator('.initial-member-option').filter(has_text='member_my').locator('input').check()
                leader.locator('.initial-member-option').filter(has_text='member_khanh').locator('input').check()
                expect(leader.locator('#initial-members-count')).to_have_text('Đã chọn 2 nhân viên')
                with leader.expect_response(lambda r:'/api/projects' in r.url and r.request.method=='POST') as created:
                    leader.locator('#modal-create-project button[type=submit]').click()
                assert created.value.status==201
                leader.wait_for_url('**/dashboard')
                pid=next(p['id'] for p in leader_ctx.request.get(base+'/api/projects').json()['projects'] if p['project_code']=='BROWSER-QA')
                expect(member.locator(f'#project-select-input option[value="{pid}"]')).to_have_count(1,timeout=12000)
                record('Create project with staff; independent member receives update',{'members':2,'project_id':pid,'sync':'within 12s'})

                # Different browser contexts share data from one server, with role checks.
                with member.expect_navigation(wait_until='domcontentloaded'):
                    member.locator('#project-select-input').select_option(str(pid))
                member.wait_for_function(f'typeof currentProjectId !== "undefined" && String(currentProjectId)==="{pid}"')
                before=member.locator('#metric-total-tasks').inner_text()
                task=leader_ctx.request.post(base+f'/api/projects/{pid}/tasks',data={'title':'Nhiệm vụ trên điện thoại','assignee_id':3,'estimated_hours':2})
                assert task.status==201,task.text()
                tid=task.json()['task_id']
                expect(member.locator('#metric-total-tasks')).to_have_text('1',timeout=12000)
                assert member_ctx.request.get(base+f'/api/projects/{pid}/tasks').json()['tasks'][0]['title']=='Nhiệm vụ trên điện thoại'
                assert admin_ctx.request.get(base+f'/api/projects/{pid}/tasks').status==200
                record('PM writes; member Dashboard and Admin read same task',{'before':before,'after':'1','admin_read':200})

                # Pending member must be visible and actionable in Admin UI.
                signup_ctx=context();signup=signup_ctx.new_page();signup.goto(base+'/login')
                signup.locator('#tab-btn-register').click()
                for id,value in {'reg-fullname':'Nhân viên đăng ký','reg-username':'browser_pending','reg-email':'browser_pending@example.test','reg-password':'QApassword123!','reg-confirm-password':'QApassword123!'}.items(): signup.locator('#'+id).fill(value)
                signup.locator('#reg-role').select_option('member')
                with signup.expect_response('**/api/auth/register') as registered: signup.locator('#btn-submit-register').click()
                assert registered.value.status==202
                assert signup_ctx.request.post(base+'/api/auth/login',data={'username_or_email':'browser_pending','password':'QApassword123!'}).status==401
                admin.bring_to_front()
                request_card=admin.locator('#pending-requests-list > div').filter(has_text='browser_pending')
                expect(request_card).to_be_visible(timeout=12000)
                request_card.get_by_role('button',name='Duyệt',exact=True).click()
                expect(request_card).to_have_count(0)
                assert signup_ctx.request.post(base+'/api/auth/login',data={'username_or_email':'browser_pending','password':'QApassword123!'}).status==200
                record('Member registration, pending Admin list, approval and login',{'registration':202,'before':401,'after':200})

                # Word upload -> extraction -> Gemini adapter -> draft displayed.
                from docx import Document
                word=Path(directory)/'bien_ban.docx';doc=Document()
                doc.add_paragraph('Cuộc họp kiểm tra tiến độ dự án.')
                doc.add_paragraph('Nhóm trao đổi khả năng sử dụng trên điện thoại và máy tính bảng.')
                doc.save(word)
                member.goto(base+'/ai-assistant');member.wait_for_function("currentAIMode==='meeting_minutes'")
                member.locator('#ai-minute-file-input').set_input_files(str(word))
                expect(member.locator('#ai-raw-notes')).to_have_value(re.compile('Cuộc họp'))
                with member.expect_response('**/api/ai') as generated: member.locator('#btn-generate-ai').click()
                assert generated.value.status==200,generated.value.text()
                expect(member.locator('#ai-output-content')).to_contain_text('Biên bản')
                assert calls and all(model == 'gemini-3.6-flash' for model in calls)
                assert member.locator('#btn-confirm-draft').is_hidden()
                record('DOCX meeting minutes through Gemini adapter',{'draft':200,'model':'gemini-3.6-flash','calls':len(calls),'member_cannot_confirm':True,'AI_HTTP':'mock'})

                # Actual UI workflow while the upstream transport is denied.
                with patch.dict(os.environ, {'AI_LOCAL_FALLBACK_ENABLED':'true','LOCAL_AI_MODEL_ENABLED':'false'}):
                    network_mode['denied']=True
                    member.locator('#ai-note-select').select_option('')
                    member.locator('#ai-raw-notes').fill('Đã thống nhất dùng SQLite.\nLong sẽ kiểm thử trước 30/10/2026.')
                    with member.expect_response('**/api/ai') as offline_minutes:
                        member.locator('#btn-generate-ai').click()
                    assert offline_minutes.value.status==200,offline_minutes.value.text()
                    expect(member.locator('#ai-engine-source')).to_contain_text('bộ xử lý nghiệp vụ cục bộ')
                    expect(member.locator('#ai-output-content')).to_contain_text('Long sẽ kiểm thử')
                    assert member.locator('#btn-confirm-draft').is_hidden()
                    member.screenshot(path=str(shots/'local-brain-meeting.png'),full_page=True)
                    record('Offline minutes UI survives API network denial',{'http':200,'engine':'local_rules','member_cannot_confirm':True})

                    leader.goto(base+'/ai-assistant')
                    leader.wait_for_function('currentUser !== null && currentProjectId !== null')
                    with leader.expect_navigation(wait_until='domcontentloaded'):
                        leader.locator('#project-select-input').select_option(str(pid))
                    leader.wait_for_function(f'typeof currentProjectId !== "undefined" && String(currentProjectId)==="{pid}"')
                    leader.goto(base+'/ai-assistant')
                    expect(leader.locator('#btn-generate-ai')).to_be_enabled()
                    leader.locator('#ai-week-start').fill('2026-10-01')
                    leader.locator('#ai-week-end').fill('2026-11-30')
                    with leader.expect_response('**/api/ai') as offline_summary:
                        leader.locator('#btn-generate-ai').click()
                    assert offline_summary.value.status==200,offline_summary.value.text()
                    expect(leader.locator('#ai-engine-source')).to_contain_text('bộ xử lý nghiệp vụ cục bộ')
                    expect(leader.locator('#ai-output-content')).to_contain_text('Tổng quan tiến độ')
                    record('Offline weekly summary UI survives API network denial',{'http':200,'engine':'local_rules'})

                    before_assignment=leader_ctx.request.get(base+f'/api/projects/{pid}/tasks').json()['tasks']
                    leader.locator('[data-mode="assignment_suggestion"]').click()
                    leader.locator('#ai-task-select').select_option(str(tid))
                    with leader.expect_response('**/api/ai') as offline_assignment:
                        leader.locator('#btn-generate-ai').click()
                    assert offline_assignment.value.status==200,offline_assignment.value.text()
                    expect(leader.locator('#ai-output-content')).to_contain_text('Gợi ý phân công')
                    expect(leader.locator('#ai-engine-source')).to_contain_text('bộ xử lý nghiệp vụ cục bộ')
                    after_assignment=leader_ctx.request.get(base+f'/api/projects/{pid}/tasks').json()['tasks']
                    assert before_assignment==after_assignment
                    record('Offline assignment UI retains existing assignments',{'http':200,'engine':'local_rules','tasks_unchanged':True})
                    network_mode['denied']=False

                member.get_by_role('button',name='Thêm ghi chú cuộc họp',exact=True).click()
                member.locator('#meeting-note-title').fill('Ghi chú lưu từ giao diện')
                member.locator('#meeting-note-attendees').fill('Lưu Trà My')
                member.locator('#meeting-note-content').fill('Cuộc họp rà soát nhiệm vụ theo tài liệu.')
                member.locator('#modal-meeting-note button[type=submit]').click()
                expect(member.locator('#modal-meeting-note')).not_to_be_visible()
                expect(member.locator('#ai-note-select option:checked')).to_contain_text('Ghi chú lưu từ giao diện')
                record('Save meeting note and select it as AI source',{'saved_and_selected':True})

                # Protect in-progress notes while another client changes the project.
                member.locator('#ai-raw-notes').fill('Ghi chú đang nhập cần được giữ nguyên')
                changed=leader_ctx.request.put(base+f'/api/tasks/{tid}',data={'title':'Thay đổi từ PM','description':'','status':'todo','priority':'medium','estimated_hours':2,'assignee_id':3})
                assert changed.status==200,changed.text()
                member.bring_to_front()
                expect(member.locator('#sync-banner')).to_be_visible(timeout=12000)
                expect(member.locator('#ai-raw-notes')).to_have_value('Ghi chú đang nhập cần được giữ nguyên')
                record('Cross-client update preserves unsaved notes',{'banner':True,'notes_preserved':True})

                for width,height,label in [(1366,900,'laptop'),(768,1024,'tablet'),(390,844,'phone'),(320,740,'small-phone')]:
                    for role,ctx in [('leader',leader_ctx),('member',member_ctx),('admin',admin_ctx)]:
                        page={'leader':leader,'member':member,'admin':admin}[role]
                        page.set_viewport_size({'width':width,'height':height})
                        paths=['/dashboard','/kanban','/sprints','/reports','/ai-assistant','/notifications']
                        if role=='admin':paths.append('/admin')
                        for path in paths:
                            page.goto(base+path,wait_until='domcontentloaded')
                            page.wait_for_function('currentUser !== null')
                            # Wait for page render, not external font/network idleness.
                            page.wait_for_timeout(180)
                            geometry=page.evaluate('({scroll:document.documentElement.scrollWidth,width:innerWidth})')
                            assert geometry['scroll']<=width+2,f'{role} {path} {width}: {geometry}'
                        if role=='leader':
                            page.get_by_role('button',name='+ Tạo dự án',exact=True).click()
                            expect(page.locator('#modal-create-project')).to_be_visible()
                            page.locator('#initial-members-search').fill('my')
                            expect(page.locator('.initial-member-option')).to_have_count(1)
                            box=page.locator('#modal-create-project .modal-box')
                            assert box.evaluate('el=>el.scrollWidth<=el.clientWidth+2')
                            box.evaluate('el=>el.scrollTop=0')
                            page.screenshot(path=str(shots/f'create-project-{label}.png'))
                            assert page.evaluate('document.documentElement.scrollWidth')<=width+2
                            page.locator('#modal-create-project .modal-header button').click()
                        if width<993:
                            page.get_by_role('button',name='Mở điều hướng',exact=True).click()
                            expect(page.locator('#app-sidebar')).to_have_class(re.compile('mobile-open'))
                            page.locator('#sidebar-backdrop').click(position={'x':max(285,width-15),'y':150},force=True)
                            expect(page.locator('#app-sidebar')).not_to_have_class(re.compile('mobile-open'))
                        record(f'Responsive {label}: {role}',{'width':width,'pages':len(paths),'body_overflow':False,'navigation':True})
                    member.goto(base+'/kanban');member.wait_for_function('currentTasks.length>0')
                    member.locator('.kanban-card, .task-card').first.locator('.task-title').click()
                    expect(member.locator('#detail-status')).to_be_visible()
                    member.locator('#detail-status').select_option('review')
                    member.get_by_role('button',name='Lưu trạng thái',exact=True).click()
                    expect(member.locator('#detail-task-status')).to_have_text('REVIEW')
                    box=member.locator('#modal-task-detail .modal-box')
                    assert box.evaluate('el=>el.scrollWidth<=el.clientWidth+2')
                    box.evaluate('el=>el.scrollTop=0')
                    member.screenshot(path=str(shots/f'kanban-{label}.png'))
                    member.locator('#modal-task-detail .modal-header button').click()
                record('Touch-friendly status change without drag and drop',{'viewports':4,'status':'review'})
                for width,height in [(390,844),(768,1024)]:
                    touch=context(width,height);page=login(touch,'member_my')
                    page.goto(base+'/kanban');page.wait_for_function('currentTasks.length>0')
                    page.locator('.kanban-card, .task-card').first.locator('.task-title').tap()
                    expect(page.locator('#detail-status')).to_be_visible()
                    page.locator('#detail-status').select_option('todo')
                    page.get_by_role('button',name='Lưu trạng thái',exact=True).tap()
                    expect(page.locator('#detail-task-status')).to_have_text('TODO')
                    touch.close()
                record('Emulated touch input on phone and tablet',{'widths':[390,768],'has_touch':True,'tap_status_saved':True})
                assert not errors,errors
                record('JavaScript runtime errors',{'errors':errors})
            except Exception as error:
                records.append({'name':'Browser acceptance','passed':False,'error':str(error)})
                for page in browser.contexts:
                    if page.pages:
                        try: page.pages[0].screenshot(path=str(shots/f'failure-{len(list(shots.glob("failure-*.png")))}.png'),full_page=True)
                        except Exception: pass
                raise
            finally:
                (out/'qa-browser-results.json').write_text(json.dumps({'browser':'Microsoft Edge headless','database':'disposable SQLite','AI':'Gemini adapter with mocked HTTP; no real key','records':records,'runtime_errors':errors,'success':all(r['passed'] for r in records)},ensure_ascii=False,indent=2),encoding='utf-8')
                browser.close();server.shutdown()
    print('Browser checks:',len(records),'passed:',sum(r['passed'] for r in records))


if __name__=='__main__':main()
