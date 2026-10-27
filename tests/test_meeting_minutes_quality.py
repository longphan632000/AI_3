import io
import json
import re
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from app import app
from database.connection import execute_db, query_db
from services.ai_confirmation_service import confirm_ai_draft
from services.meeting_minutes_service import create_meeting_minutes
from services.gemini_service import AIProviderError
from services.project_data_retriever import retrieve_ai_data


class FakeClassifier:
    def __init__(self, omit_last=False):
        self.calls = 0
        self.omit_last = omit_last

    def generate_response(self, prompt, max_output_tokens=4096, json_output=False):
        self.calls += 1
        source = prompt.split('NGUỒN:\n', 1)[1]
        items = []
        for line in source.splitlines():
            match = re.match(r'^\[(\d+)\] (.*)', line)
            if not match:
                continue
            identifier, text = int(match.group(1)), match.group(2)
            kinds = []
            if 'quyết định' in text.lower():
                kinds.append('decision')
            if 'giao' in text.lower():
                kinds.append('action')
            if not kinds:
                kinds = ['discussion']
            items.append({'id': identifier, 'kinds': kinds, 'summary': text[:100]})
        if self.omit_last:
            items.pop()
        return json.dumps({'items': items}, ensure_ascii=False)


class TestMeetingMinutesQuality(unittest.TestCase):
    def setUp(self):
        self.client = app.test_client()
        self.client.post('/api/auth/login', json={
            'username_or_email': 'leader_long', 'password': 'Password123!'
        })

    def test_long_minutes_cover_all_passages_and_store_actual_fields(self):
        notes = '\n'.join(f'Thảo luận hạng mục {i} và rủi ro {i}.' for i in range(90))
        notes += '\nQuyết định dùng phương án A và giao Lan hoàn tất trước 30/09.'
        fake = FakeClassifier()
        content, structured, refs, warnings = create_meeting_minutes(None, notes, fake)
        self.assertGreater(fake.calls, 1)
        self.assertEqual(len(refs), 91)
        self.assertIn('hạng mục 89', content)
        self.assertIn('phương án A', structured['decisions'])
        self.assertIn('giao Lan', structured['action_items'])
        self.assertFalse(warnings)

    def test_existing_note_and_uploaded_text_are_both_processed(self):
        note = {'id': 7, 'title': 'Họp kế hoạch', 'meeting_date': '2026-09-25',
                'attendees': 'Lan, Minh', 'content': 'Quyết định chọn A.'}
        content, structured, refs, warnings = create_meeting_minutes(
            note, 'Giao Lan kiểm thử trước 30/09.', FakeClassifier()
        )
        self.assertEqual(len(refs), 2)
        self.assertEqual(structured['meeting_note_id'], 7)
        self.assertIn('Quyết định chọn A', structured['decisions'])
        self.assertIn('Giao Lan', structured['action_items'])
        self.assertFalse(warnings)

    def test_action_owner_and_deadline_must_appear_in_source_quote(self):
        class ActionClassifier:
            def __init__(self, owner):
                self.owner = owner

            def generate_response(self, prompt, **kwargs):
                return json.dumps({'items': [{
                    'id': 1, 'kinds': ['action'], 'summary': 'Giao việc kiểm thử.',
                    'actions': [{'quote': 'Giao Lan kiểm thử trước 30/09.',
                                 'owner': self.owner, 'deadline': '30/09'}]
                }]}, ensure_ascii=False)

        note = 'Giao Lan kiểm thử trước 30/09.'
        _, valid, _, warnings = create_meeting_minutes(None, note, ActionClassifier('Lan'))
        self.assertEqual(valid['action_details'][0]['owner'], 'Lan')
        self.assertEqual(valid['action_details'][0]['deadline'], '30/09')
        self.assertFalse(warnings)
        _, invalid, _, warnings = create_meeting_minutes(None, note, ActionClassifier('Minh'))
        self.assertEqual(invalid['action_details'][0]['owner'], '')
        self.assertTrue(warnings)

    def test_explicit_decisions_and_assignment_table_keep_source_fields(self):
        notes = ('4. Quyết định đã thống nhất\n'
                 'Chọn phương án A.\n'
                 '5. Công việc sau cuộc họp\n'
                 'STT | Công việc | Phụ trách | Hạn hoàn thành | Ưu tiên | Trạng thái\n'
                 '1 | Kiểm thử phương án A | Lan | 30/09/2026 | Cao | To Do')
        _, structured, _, warnings = create_meeting_minutes(None, notes, FakeClassifier())
        self.assertIn('Chọn phương án A.', structured['decisions'])
        self.assertIn('Kiểm thử phương án A', structured['action_items'])
        self.assertEqual(structured['action_details'][0]['owner'], 'Lan')
        self.assertEqual(structured['action_details'][0]['deadline'], '30/09/2026')
        self.assertFalse(warnings)

    def test_missing_classification_is_retried_only_for_missing_passage(self):
        class OmitsOnce(FakeClassifier):
            def generate_response(self, prompt, **kwargs):
                self.calls += 1
                source = prompt.split('NGUỒN:\n', 1)[1]
                ids = [int(match.group(1)) for line in source.splitlines()
                       if (match := re.match(r'^\[(\d+)\] ', line))]
                if self.calls == 1:
                    ids.pop()
                return json.dumps({'items': [
                    {'id': identifier, 'kinds': ['discussion'], 'summary': 'Nội dung thảo luận.'}
                    for identifier in ids
                ]})

        provider = OmitsOnce()
        content, _, refs, warnings = create_meeting_minutes(
            None, 'Nội dung thảo luận.\nNội dung thảo luận bổ sung.', provider
        )
        self.assertEqual(provider.calls, 2)
        self.assertEqual(len(refs), 2)
        self.assertIn('bổ sung', content)
        self.assertFalse(warnings)

    def test_unsupported_summary_number_is_discarded_in_favor_of_source(self):
        class InventedNumber:
            def generate_response(self, prompt, **kwargs):
                return json.dumps({'items': [
                    {'id': 1, 'kinds': ['discussion'], 'summary': 'Rủi ro số 999.'}
                ]})

        content, _, _, warnings = create_meeting_minutes(
            None, 'Rủi ro cần kiểm tra thêm.', InventedNumber()
        )
        self.assertIn('Rủi ro cần kiểm tra thêm.', content)
        self.assertNotIn('999', content)
        self.assertFalse(warnings)

    def test_missing_passage_is_visible_and_cannot_be_confirmed(self):
        with patch('services.ai_provider_service.AIProviderService.generate_response',
                   side_effect=FakeClassifier(omit_last=True).generate_response):
            response = self.client.post('/api/ai', json={
                'task_type': 'meeting_minutes', 'project_id': 1,
                'raw_notes': 'Quyết định chọn A.\nGiao Lan làm trước thứ Sáu.'
            })
        self.assertEqual(response.status_code, 200)
        result = response.get_json()
        self.assertEqual(result['status'], 'needs_verification')
        self.assertIn('Giao Lan', result['result'])
        self.assertTrue(result['warnings'])
        confirm = self.client.post(f"/api/ai/drafts/{result['draft_id']}/confirm")
        self.assertEqual(confirm.status_code, 403)

    def test_complete_minutes_confirm_exact_decisions_and_actions(self):
        with patch('services.ai_provider_service.AIProviderService.generate_response',
                   side_effect=FakeClassifier().generate_response):
            response = self.client.post('/api/ai', json={
                'task_type': 'meeting_minutes', 'project_id': 1,
                'raw_notes': 'Quyết định dùng phương án xanh 2048.\nGiao Lan kiểm thử trước 30/09.'
            })
        self.assertEqual(response.status_code, 200)
        result = response.get_json()
        self.assertEqual(result['status'], 'draft')
        confirmed = self.client.post(f"/api/ai/drafts/{result['draft_id']}/confirm")
        self.assertEqual(confirmed.status_code, 200)
        minute = query_db('''
            SELECT decisions, action_items FROM meeting_minutes
            WHERE summary LIKE %s ORDER BY id DESC LIMIT 1
        ''', ('%phương án xanh 2048%',), one=True)
        self.assertEqual(minute['decisions'], 'Quyết định dùng phương án xanh 2048.')
        self.assertEqual(minute['action_items'], 'Giao Lan kiểm thử trước 30/09.')
        history = query_db('SELECT confirmation_status FROM ai_history WHERE draft_id = %s',
                           (result['draft_id'],), one=True)
        self.assertEqual(history['confirmation_status'], 'confirmed')

    def test_docx_preserves_paragraph_table_order(self):
        from docx import Document
        document = Document()
        document.add_paragraph('Mở đầu')
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = 'Quyết định'
        table.cell(0, 1).text = 'Phương án A'
        document.add_paragraph('Kết thúc')
        stream = io.BytesIO()
        document.save(stream)
        stream.seek(0)
        response = self.client.post('/api/ai/upload-meeting-file',
                                    data={'file': (stream, 'minutes.docx')},
                                    content_type='multipart/form-data')
        self.assertEqual(response.status_code, 200)
        text = response.get_json()['extracted_text']
        self.assertLess(text.index('Mở đầu'), text.index('Quyết định'))
        self.assertLess(text.index('Quyết định'), text.index('Kết thúc'))

    def test_unreadable_pdf_returns_clear_error(self):
        from pypdf import PdfWriter
        writer = PdfWriter()
        writer.add_blank_page(width=200, height=200)
        stream = io.BytesIO()
        writer.write(stream)
        stream.seek(0)
        response = self.client.post('/api/ai/upload-meeting-file',
                                    data={'file': (stream, 'scanned.pdf')},
                                    content_type='multipart/form-data')
        self.assertEqual(response.status_code, 400)
        self.assertIn('không chứa nội dung', response.get_json()['error'])

    def test_weekly_retrieval_uses_requested_dates(self):
        task_id = execute_db('''
            INSERT INTO tasks (project_id, reporter_id, title, status, priority, due_date, created_at, updated_at)
            VALUES (1, 2, 'Kiểm thử lọc tuần', 'todo', 'medium', '2026-09-15', '2026-09-10', '2026-09-10')
        ''')
        request = {'task_type': 'weekly_summary', 'project_id': 1,
                   'week_start': '2026-09-01', 'week_end': '2026-09-30'}
        included, _ = retrieve_ai_data(request)
        self.assertIn(task_id, [task['task_id'] for task in included['tasks']])
        request.update(week_start='2025-01-01', week_end='2025-01-31')
        excluded, _ = retrieve_ai_data(request)
        self.assertNotIn(task_id, [task['task_id'] for task in excluded['tasks']])

    def test_missing_provider_key_creates_no_draft(self):
        before = query_db('SELECT COUNT(*) AS total FROM ai_drafts', one=True)['total']
        with tempfile.TemporaryDirectory() as directory, \
                patch.dict('os.environ', {'GEMINI_API_KEY': ''}), \
                patch('services.ai_config.ENV_FILE', Path(directory) / 'missing.env'):
            response = self.client.post('/api/ai', json={
                'task_type': 'meeting_minutes', 'project_id': 1,
                'raw_notes': 'Quyết định chọn phương án A.'
            })
        self.assertEqual(response.status_code, 503)
        after = query_db('SELECT COUNT(*) AS total FROM ai_drafts', one=True)['total']
        self.assertEqual(before, after)

    def test_failed_word_chunk_reports_position_and_saves_no_partial_draft(self):
        before = query_db('SELECT COUNT(*) AS total FROM ai_drafts', one=True)['total']
        notes = '\n'.join(f'Nội dung cuộc họp đoạn {number}.' for number in range(45))
        with patch('services.ai_provider_service.AIProviderService.generate_response',
                   side_effect=AIProviderError('Gemini xử lý quá thời gian chờ.', 'AI_TIMEOUT')):
            response = self.client.post('/api/ai', json={
                'task_type': 'meeting_minutes', 'project_id': 1, 'raw_notes': notes
            })
        self.assertEqual(response.status_code, 503)
        self.assertEqual(response.get_json()['error_code'], 'AI_TIMEOUT')
        self.assertIn('phần 1/3', response.get_json()['error'])
        after = query_db('SELECT COUNT(*) AS total FROM ai_drafts', one=True)['total']
        self.assertEqual(before, after)


if __name__ == '__main__':
    unittest.main()
